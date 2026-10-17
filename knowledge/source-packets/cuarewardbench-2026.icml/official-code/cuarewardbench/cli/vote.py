#!/usr/bin/env python3
"""
voting_rm.py
对多个 RM 结果文件进行投票后重新计算指标。
"""

import json
import argparse
import os
from collections import defaultdict
from cuarewardbench.core.utils import compute_and_save_metrics, parse_action_reward, compute_binary_classification_metrics, analyze_group_performance
from cuarewardbench.core.config import get_vote_config

# 添加Excel导出相关的导入
try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:  # optional unless Excel export is required
    Workbook = Font = Alignment = Border = Side = get_column_letter = None

parser = argparse.ArgumentParser(description="Vote among multiple RM detailed results and re-compute metrics")
parser.add_argument("--output_dir", default="./",
                    help="Directory to save metrics file")
parser.add_argument("--reward_type", default="or_ar",
                    choices=["or", "ar", "or_ar"],
                    help="Which reward metrics to compute (or, ar, or both)")
parser.add_argument("--vote_mode", type=str, choices=['all', 'one', 'majority', 'upe', 'majority_dual'], default='upe',
                    help="投票模式: all: 所有评审者一致; one: 至少一个评审者同意; majority: 大多数评审者同意; upe: 全票一致才输出，否则None; majority_dual: 过半一致才输出，否则None")
parser.add_argument("--vote_config", type=str, default="upe_orm_2_qwen3",
                    help="投票配置: 指定预设详细结果文件列表")
parser.add_argument("--input_files", nargs="+", default=None,
                    help="显式指定参与投票的 detailed_results JSON 文件；传入后优先于 --vote_config")
args = parser.parse_args()

# 支持两种输入方式：显式 --input_files 优先；否则从白名单 vote_config 读取。
if args.input_files:
    detailed_filepath_list = args.input_files
else:
    detailed_filepath_list = get_vote_config(args.vote_config)

# 加载所有 RM 的详细结果
all_rm_results = []
for fp in detailed_filepath_list:
    if not os.path.exists(fp):
        print(f"[Warn] 结果文件不存在: {fp}")
        continue
    with open(fp, "r", encoding="utf-8") as f:
        all_rm_results.append(json.load(f))
if not all_rm_results:
    raise RuntimeError("未找到任何有效结果文件！")

# ---------- 投票合并 ----------
def vote(predictions, mode, rm_count):
    """根据投票模式返回最终预测值，支持返回None"""
    positive_votes = sum(predictions)
    
    # 添加调试信息：当只有两个模型时，检查upe和majority_dual是否等价
    if rm_count == 2:
        # 计算两种模式的结果
        upe_result = None
        if positive_votes == rm_count:  # 全票赞成
            upe_result = 1
        elif positive_votes == 0:       # 全票反对
            upe_result = 0
        else:                           # 无法达成共识
            upe_result = None
            
        majority_dual_result = None
        if positive_votes > rm_count / 2:     # 过半赞成
            majority_dual_result = 1
        elif positive_votes < rm_count / 2:   # 过半反对
            majority_dual_result = 0
        else:                                 # 平票，无法达成共识
            majority_dual_result = None
            
        # 检查是否不等价
        if upe_result != majority_dual_result:
            print(f"[DEBUG] 检测到upe和majority_dual结果不一致！")
            print(f"  predictions: {predictions}")
            print(f"  positive_votes: {positive_votes}, rm_count: {rm_count}")
            print(f"  upe_result: {upe_result}")
            print(f"  majority_dual_result: {majority_dual_result}")
    
    # 原有的投票逻辑
    if mode == 'all':          # 全票赞成
        return 1 if positive_votes == rm_count else 0
    elif mode == 'one':        # ≥ 1 票
        return 1 if positive_votes > 0 else 0
    elif mode == 'majority':        # 过半
        return 1 if positive_votes > rm_count / 2 else 0
    elif mode == 'upe':   # 全票一致才输出，否则None
        if positive_votes == rm_count:  # 全票赞成
            return 1
        elif positive_votes == 0:       # 全票反对
            return 0
        else:                             # 无法达成共识
            return None
    elif mode == 'majority_dual':  # 过半一致才输出，否则None
        if positive_votes > rm_count / 2:     # 过半赞成
            return 1
        elif positive_votes < rm_count / 2:   # 过半反对
            return 0
        else:                                 # 平票，无法达成共识
            return None
    else:
        raise ValueError("未知的 vote_mode")

def compute_metrics_with_none(stats):
    """计算包含None样本的指标"""
    tp, fp, fn, tn = stats['tp'], stats['fp'], stats['fn'], stats['tn']
    none_count = stats.get('none', 0)
    none_positive = stats.get('none_positive', 0)  # 真实为正但无法预测
    none_negative = stats.get('none_negative', 0)  # 真实为负但无法预测
    
    # Precision：不受None影响（None没有被预测为正）
    if tp + fp > 0:
        precision = tp / (tp + fp)
    else:
        precision = 0.0
    
    # Recall：分母包含真实为正但无法预测的样本
    if tp + fn + none_positive > 0:
        recall = tp / (tp + fn + none_positive)
    else:
        recall = 0.0
    
    # Specificity：分母包含真实为负但无法预测的样本
    if tn + fp + none_negative > 0:
        specificity = tn / (tn + fp + none_negative)
    else:
        specificity = 0.0
    
    # NPV：不受None影响（None没有被预测为负）
    if tn + fn > 0:
        npv = tn / (tn + fn)
    else:
        npv = 0.0
    
    # F1 score
    if precision + recall > 0:
        f1 = 2 * precision * recall / (precision + recall)
    else:
        f1 = 0.0
    
    # Overall Accuracy：包含None作为错误预测
    total_samples = tp + fp + fn + tn + none_count
    if total_samples > 0:
        oa_with_none = (tp + tn) / total_samples
        coverage = (tp + fp + fn + tn) / total_samples
    else:
        oa_with_none = 0.0
        coverage = 0.0
    
    # 传统指标（仅针对可预测样本）
    predictable_samples = tp + fp + fn + tn
    if predictable_samples > 0:
        oa_predictable = (tp + tn) / predictable_samples
    else:
        oa_predictable = 0.0
    
    return {
        'tp': tp,
        'fp': fp,
        'fn': fn,
        'tn': tn,
        'precision': precision,
        'recall': recall,
        'specificity': specificity,  # 新增specificity指标
        'npv': npv,
        'f1': f1,
        'oa_with_none': oa_with_none,      # 包含None的总体准确率
        'oa_predictable': oa_predictable,   # 仅可预测样本的准确率
        'coverage': coverage,               # 模型覆盖率
        'none_count': none_count,
        'none_positive': none_positive,
        'none_negative': none_negative,
        'total_samples': total_samples,
        'predictable_samples': predictable_samples,
        # 保持向后兼容
        'accuracy': oa_with_none,  # 默认使用包含None的准确率
        'total': predictable_samples,
        'fail': 0  # 保持兼容性
    }

def collect_statistics_with_none(all_results, stats, task_type_stats):
    """收集统计信息，支持None预测值"""
    for task in all_results:
        pred = task["pred"]
        true_label = task["gt"]  # 修复：使用gt字段而不是label字段
        model_setting = task["model_setting"]
        task_type = task["task_type"]
        
        # 更新overall统计
        stats["overall"]["total"] += 1
        task_type_stats[task_type]["total"] += 1
        stats[model_setting]["total"] += 1
        
        if pred is None:
            # None样本统计
            stats["overall"]["none"] = stats["overall"].get("none", 0) + 1
            task_type_stats[task_type]["none"] = task_type_stats[task_type].get("none", 0) + 1
            stats[model_setting]["none"] = stats[model_setting].get("none", 0) + 1
            
            if true_label == 1:
                stats["overall"]["none_positive"] = stats["overall"].get("none_positive", 0) + 1
                task_type_stats[task_type]["none_positive"] = task_type_stats[task_type].get("none_positive", 0) + 1
                stats[model_setting]["none_positive"] = stats[model_setting].get("none_positive", 0) + 1
            else:
                stats["overall"]["none_negative"] = stats["overall"].get("none_negative", 0) + 1
                task_type_stats[task_type]["none_negative"] = task_type_stats[task_type].get("none_negative", 0) + 1
                stats[model_setting]["none_negative"] = stats[model_setting].get("none_negative", 0) + 1
        else:
            # 正常预测样本统计
            if pred == 1 and true_label == 1:
                stats["overall"]["tp"] += 1
                task_type_stats[task_type]["tp"] += 1
                stats[model_setting]["tp"] += 1
            elif pred == 1 and true_label == 0:
                stats["overall"]["fp"] += 1
                task_type_stats[task_type]["fp"] += 1
                stats[model_setting]["fp"] += 1
            elif pred == 0 and true_label == 1:
                stats["overall"]["fn"] += 1
                task_type_stats[task_type]["fn"] += 1
                stats[model_setting]["fn"] += 1
            elif pred == 0 and true_label == 0:
                stats["overall"]["tn"] += 1
                task_type_stats[task_type]["tn"] += 1
                stats[model_setting]["tn"] += 1

# ---------- 用 task_id + model_setting 作为 key 对齐所有 RM 结果 ----------
# 构建索引：{ (task_id, model_setting): task_dict } 形式
rm_indexed = []
for rm_res in all_rm_results:
    idx_map = {(item["task_id"], item["model_setting"]): item for item in rm_res}
    rm_indexed.append(idx_map)

# 以第一个 RM 为基准，逐个任务投票
merged_results = []
none_count = 0

# 只有在需要计算轨迹级指标时才进行轨迹级投票
if "or" in args.reward_type:
    for (task_id, model_setting), task_base in rm_indexed[0].items():
        preds = []
        for idx_map in rm_indexed:
            if (task_id, model_setting) not in idx_map:
                raise KeyError(f"缺失任务: {task_id} - {model_setting}")
            preds.append(idx_map[(task_id, model_setting)]["pred"])
        voted_pred = vote(preds, args.vote_mode, len(all_rm_results))

        if voted_pred is None:
            none_count += 1

        voted_task = task_base.copy()
        voted_task["pred"] = voted_pred
        merged_results.append(voted_task)

    print(f"轨迹级投票完成，共 {len(merged_results)} 个样本，其中 {none_count} 个无法达成共识（None）")
else:
    # 如果只计算action reward，直接使用第一个RM的结果作为基准（不进行轨迹级投票）
    merged_results = list(rm_indexed[0].values())
    print(f"跳过轨迹级投票（仅计算action reward），共 {len(merged_results)} 个样本")

# ---------- 轨迹级统计 ----------
stats = defaultdict(lambda: {"tp": 0, "fp": 0, "fn": 0, "tn": 0, "fail": 0, "total": 0, "none": 0, "none_positive": 0, "none_negative": 0})
task_type_stats = defaultdict(lambda: {"tp": 0, "fp": 0, "fn": 0, "tn": 0, "fail": 0, "total": 0, "none": 0, "none_positive": 0, "none_negative": 0})

# ---------- 改写 action-level 的解析函数 ----------
def voting_parse_action_reward(all_rm_results, vote_mode):
    """action-level 投票解析：使用 (task_id, model_setting) 对齐后投票"""
    # 对每个 RM 的结果进行 parse_action_reward 解析
    rm_parsed = [parse_action_reward(rm)[0] for rm in all_rm_results]

    # 构建索引：{ (task_id, model_setting): task_dict }
    rm_indexed = []
    for parsed in rm_parsed:
        idx_map = {(item["task_id"], item["model_setting"]): item for item in parsed}
        rm_indexed.append(idx_map)

    # 以第一个 RM 为基准逐个任务投票
    base_index = rm_indexed[0]
    merged_tasks = []
    all_true_labels, all_pred_labels = [], []

    for key, base_task in base_index.items():
        # 收集所有 RM 的 action_reward_list
        action_reward_lists = []
        for idx_map in rm_indexed:
            if key not in idx_map:
                raise KeyError(f"缺失任务: {key}")
            action_reward_lists.append(idx_map[key]["action_reward_list"])

        # 逐 step 投票
        step_num = base_task["step_num"]
        voted_action_reward_list = []
        has_none = False
        
        for i in range(step_num):
            step_vote = vote([lst[i] for lst in action_reward_lists], vote_mode, len(all_rm_results))
            voted_action_reward_list.append(step_vote)
            if step_vote is None:
                has_none = True

        # 构造 true/pred 标签 - 不再跳过包含None的task
        action_annos = base_task.get("action_annos")
        if action_annos:
            good_steps = (
                action_annos.get('key_good_before_bad_in_fail', []) +
                action_annos.get('key_good_in_success', []) +
                action_annos.get('key_good_after_bad_in_success', [])
            )
            bad_steps = (
                action_annos.get('key_bad_in_fail', []) +
                action_annos.get('key_bad_in_success', [])
            )

            for s in good_steps:
                if 1 <= s <= step_num:
                    all_true_labels.append(1)
                    all_pred_labels.append(voted_action_reward_list[s - 1])  # 可能是None
            for s in bad_steps:
                if 1 <= s <= step_num:
                    all_true_labels.append(0)
                    all_pred_labels.append(voted_action_reward_list[s - 1])  # 可能是None

        # 回填新的 action_reward_list
        voted_task = base_task.copy()
        voted_task["action_reward_list"] = voted_action_reward_list
        merged_tasks.append(voted_task)

    return merged_tasks, all_true_labels, all_pred_labels

def extract_metrics_for_excel(metrics_data):
    """从指标数据中提取Excel所需的格式化数据"""
    # 定义任务类型映射（JSON中的键名到表格列名的映射）
    task_mapping = {
        'vs_code': 'vscode',
        'gimp': 'gimp', 
        'libreoffice_writer': 'writer',
        'chrome': 'chrome',
        'multi_apps': 'multi_apps',
        'vlc': 'vlc',
        'os': 'os',
        'thunderbird': 'thunderbird',
        'libreoffice_impress': 'impress',
        'libreoffice_calc': 'calc'
    }
    
    result = {}
    
    # 处理 trajectory_reward_metrics
    if 'trajectory_reward_metrics' in metrics_data:
        traj_metrics = metrics_data['trajectory_reward_metrics']
        overall_metrics = traj_metrics.get('overall', {})
        by_task_type = traj_metrics.get('by_task_type', {})
        
        traj_result = {}
        
        # Overall数据 - 添加S列
        traj_result['overall'] = {
            'P': round(overall_metrics.get('precision', 0) * 100, 1),
            'NPV': round(overall_metrics.get('npv', 0) * 100, 1),
            'R': round(overall_metrics.get('recall', 0) * 100, 1),
            'S': round(overall_metrics.get('specificity', 0) * 100, 1),  # 新增S列
            'OA': round(overall_metrics.get('accuracy', 0) * 100, 1)  # trajectory用accuracy
        }
        
        # 各任务类型数据 - 保持原有的P、NPV、R、OA
        for json_key, table_key in task_mapping.items():
            if json_key in by_task_type:
                metrics = by_task_type[json_key]
                traj_result[table_key] = {
                    'P': round(metrics.get('precision', 0) * 100, 1),
                    'NPV': round(metrics.get('npv', 0) * 100, 1),
                    'R': round(metrics.get('recall', 0) * 100, 1),
                    'OA': round(metrics.get('accuracy', 0) * 100, 1)  # trajectory用accuracy
                }
            else:
                # 如果没有数据，填充空值
                traj_result[table_key] = {'P': '', 'NPV': '', 'R': '', 'OA': ''}
        
        result['trajectory'] = traj_result
    
    # 处理 action_reward_metrics
    if 'action_reward_metrics' in metrics_data:
        action_metrics = metrics_data['action_reward_metrics']
        overall_metrics = action_metrics.get('overall', {})
        by_task_type = action_metrics.get('by_task_type', {})
        
        action_result = {}
        
        # Overall数据 - 添加S列
        action_result['overall'] = {
            'P': round(overall_metrics.get('precision', 0) * 100, 1),
            'NPV': round(overall_metrics.get('npv', 0) * 100, 1),
            'R': round(overall_metrics.get('recall', 0) * 100, 1),
            'S': round(overall_metrics.get('specificity', 0) * 100, 1),  # 新增S列
            'OA': round(overall_metrics.get('oa', 0) * 100, 1)  # action用oa
        }
        
        # 各任务类型数据 - 保持原有的P、NPV、R、OA
        for json_key, table_key in task_mapping.items():
            if json_key in by_task_type:
                metrics = by_task_type[json_key].get('metrics', {})
                action_result[table_key] = {
                    'P': round(metrics.get('precision', 0) * 100, 1),
                    'NPV': round(metrics.get('npv', 0) * 100, 1),
                    'R': round(metrics.get('recall', 0) * 100, 1),
                    'OA': round(metrics.get('oa', 0) * 100, 1)  # action用oa
                }
            else:
                # 如果没有数据，填充空值
                action_result[table_key] = {'P': '', 'NPV': '', 'R': '', 'OA': ''}
        
        result['action'] = action_result
    
    return result


def create_excel_from_metrics(metrics_data, model_name, output_file):
    """根据指标数据创建Excel文件"""
    if Workbook is None:
        print("openpyxl is not installed; skip Excel export.")
        return
    
    # 提取Excel格式的数据
    excel_data = extract_metrics_for_excel(metrics_data)
    
    # 创建工作簿
    wb = Workbook()
    ws = wb.active
    ws.title = "Reward Model Metrics"
    
    # 定义列标题
    columns = ['overall', 'vscode', 'gimp', 'writer', 'chrome', 'multi_apps', 
               'vlc', 'os', 'thunderbird', 'impress', 'calc']
    
    # 设置表头
    # 第一行：主标题
    ws.merge_cells('A1:A2')
    ws['A1'] = 'Reward Model'
    ws['A1'].font = Font(bold=True)
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    
    # 第二行：各任务类型标题
    col_start = 2
    for i, col_name in enumerate(columns):
        if i == 0:  # overall列有5列（P, NPV, R, S, OA）
            col_count = 5
            metrics = ['P', 'NPV', 'R', 'S', 'OA']
        else:  # 其他列有4列（P, NPV, R, OA）
            col_count = 4
            metrics = ['P', 'NPV', 'R', 'OA']
        
        # 合并列作为一个任务类型的标题
        start_col = col_start
        end_col = col_start + col_count - 1
        ws.merge_cells(start_row=1, start_column=start_col, end_row=1, end_column=end_col)
        cell = ws.cell(row=1, column=start_col)
        cell.value = col_name.replace('_', ' ').title()
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal='center', vertical='center')
        
        # 第二行：指标名称
        for j, metric in enumerate(metrics):
            cell = ws.cell(row=2, column=start_col + j)
            cell.value = metric
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal='center', vertical='center')
        
        col_start += col_count
    
    # 添加数据行
    current_row = 3
    
    # 添加trajectory数据行（如果存在）
    if 'trajectory' in excel_data:
        ws.cell(row=current_row, column=1).value = f"{model_name}_trajectory"
        ws.cell(row=current_row, column=1).font = Font(bold=True)
        
        col = 2
        traj_metrics = excel_data['trajectory']
        for i, col_name in enumerate(columns):
            if col_name in traj_metrics:
                data = traj_metrics[col_name]
                if i == 0:  # overall列
                    ws.cell(row=current_row, column=col).value = data['P']
                    ws.cell(row=current_row, column=col+1).value = data['NPV']
                    ws.cell(row=current_row, column=col+2).value = data['R']
                    ws.cell(row=current_row, column=col+3).value = data['S']
                    ws.cell(row=current_row, column=col+4).value = data['OA']
                    col += 5
                else:  # 其他列
                    ws.cell(row=current_row, column=col).value = data['P']
                    ws.cell(row=current_row, column=col+1).value = data['NPV']
                    ws.cell(row=current_row, column=col+2).value = data['R']
                    ws.cell(row=current_row, column=col+3).value = data['OA']
                    col += 4
            else:
                # 根据列类型跳过相应的列数
                col += 5 if i == 0 else 4
        current_row += 1
    
    # 添加action数据行（如果存在）
    if 'action' in excel_data:
        ws.cell(row=current_row, column=1).value = f"{model_name}_action"
        ws.cell(row=current_row, column=1).font = Font(bold=True)
        
        col = 2
        action_metrics = excel_data['action']
        for i, col_name in enumerate(columns):
            if col_name in action_metrics:
                data = action_metrics[col_name]
                if i == 0:  # overall列
                    ws.cell(row=current_row, column=col).value = data['P']
                    ws.cell(row=current_row, column=col+1).value = data['NPV']
                    ws.cell(row=current_row, column=col+2).value = data['R']
                    ws.cell(row=current_row, column=col+3).value = data['S']
                    ws.cell(row=current_row, column=col+4).value = data['OA']
                    col += 5
                else:  # 其他列
                    ws.cell(row=current_row, column=col).value = data['P']
                    ws.cell(row=current_row, column=col+1).value = data['NPV']
                    ws.cell(row=current_row, column=col+2).value = data['R']
                    ws.cell(row=current_row, column=col+3).value = data['OA']
                    col += 4
            else:
                # 根据列类型跳过相应的列数
                col += 5 if i == 0 else 4
        current_row += 1
    
    # 设置边框
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # 计算最大列数：overall(5列) + 其他10个任务类型(4列*10) = 45列
    max_row = current_row - 1
    max_col = 2 + 5 + 10 * 4 - 1  # A列(1) + overall(5) + 其他10个任务类型(40) = 46列
    
    # 应用边框到所有使用的单元格
    for row in range(1, max_row + 1):
        for col in range(1, max_col + 1):
            ws.cell(row=row, column=col).border = thin_border
    
    # 调整列宽
    ws.column_dimensions['A'].width = 25  # 增加第一列宽度以容纳更长的名称
    for col in range(2, max_col + 1):
        column_letter = get_column_letter(col)
        ws.column_dimensions[column_letter].width = 8
    
    # 保存文件
    wb.save(output_file)
    print(f"Excel文件已保存到: {output_file}")

# 新增支持None的二分类指标计算函数
def compute_binary_classification_metrics_with_none(true_labels, pred_labels):
    """计算包含None预测值的二分类指标"""
    if len(true_labels) != len(pred_labels):
        raise ValueError("true_labels和pred_labels长度不匹配")
    
    # 统计各种情况
    tp = fp = fn = tn = 0
    none_count = 0
    none_positive = 0  # 真实为正但无法预测
    none_negative = 0  # 真实为负但无法预测
    
    for true_label, pred_label in zip(true_labels, pred_labels):
        if pred_label is None:
            none_count += 1
            if true_label == 1:
                none_positive += 1
            else:
                none_negative += 1
        elif pred_label == 1 and true_label == 1:
            tp += 1
        elif pred_label == 1 and true_label == 0:
            fp += 1
        elif pred_label == 0 and true_label == 1:
            fn += 1
        elif pred_label == 0 and true_label == 0:
            tn += 1
    
    # 使用与trajectory相同的指标计算逻辑
    stats = {
        'tp': tp, 'fp': fp, 'fn': fn, 'tn': tn,
        'none': none_count,
        'none_positive': none_positive,
        'none_negative': none_negative
    }
    
    # 计算指标并添加oa字段以保持兼容性
    metrics = compute_metrics_with_none(stats)
    # 为action reward指标添加oa字段（使用oa_with_none的值）
    metrics['oa'] = metrics['oa_with_none']
    
    return metrics

# 新增支持None的组性能分析函数
def analyze_group_performance_with_none(all_results):
    """分析组性能，支持None预测值"""
    from cuarewardbench.core.utils import analyze_group_performance
    
    # 先过滤掉包含None的action步骤，使用原有的分析逻辑
    filtered_results = []
    for task in all_results:
        action_reward_list = task.get('action_reward_list', [])
        if None not in action_reward_list:
            filtered_results.append(task)
    
    # 对可预测的样本使用原有逻辑
    if filtered_results:
        predictable_metrics = analyze_group_performance(filtered_results)
        
        # 添加覆盖率信息
        total_tasks = len(all_results)
        predictable_tasks = len(filtered_results)
        coverage = predictable_tasks / total_tasks if total_tasks > 0 else 0
        
        # 为每个组添加覆盖率信息
        for group_name, group_metrics in predictable_metrics.items():
            group_metrics['coverage'] = coverage
            group_metrics['total_with_none'] = total_tasks
            group_metrics['predictable_total'] = predictable_tasks
        
        return predictable_metrics
    else:
        return {}

# ---------- 计算并保存 ----------
# 重写 compute_and_save_metrics 内部逻辑以支持投票
def compute_and_save_voting_metrics(all_results, stats, task_type_stats,
                                    metrics_filename, args, all_rm_results):
    """在投票结果上计算指标"""
    # 1. 轨迹级指标
    from cuarewardbench.core.utils import compute_step_num_metrics
    if "or" in args.reward_type:
        collect_statistics_with_none(all_results, stats, task_type_stats)

    # 使用新的指标计算函数
    for key in list(stats.keys()):
        if stats[key]['total'] > 0:
            stats[key] = compute_metrics_with_none(stats[key])

    for task_type in list(task_type_stats.keys()):
        task_type_stats[task_type] = compute_metrics_with_none(task_type_stats[task_type])

    # 2. 重构统计结构
    new_stats = {}
    
    # 构建trajectory_reward_metrics
    if "or" in args.reward_type:
        new_stats['trajectory_reward_metrics'] = {
            'overall': stats.get('overall', {}),
            'by_task_type': task_type_stats,
            'by_model_setting': {
                k: v for k, v in stats.items() if k != 'overall'
            },
            'by_step_num': compute_step_num_metrics(all_results)
        }

    # 3. action-level 指标 - 修改为支持None值
    if 'ar' in args.reward_type:
        # 使用投票版本的action reward解析
        all_results, all_true_labels, all_pred_labels = voting_parse_action_reward(
            all_rm_results, args.vote_mode
        )
        
        # 使用支持None的指标计算函数
        overall_action_metrics = compute_binary_classification_metrics_with_none(
            all_true_labels, all_pred_labels
        )
        
        # 按task_type分组收集标签数据
        task_type_labels = {}
        for task in all_results:
            task_type = task['task_type']
            if task_type not in task_type_labels:
                task_type_labels[task_type] = {'true_labels': [], 'pred_labels': []}
            
            # 从已经处理好的数据中提取标签
            action_annos = task.get('action_annos', {})
            action_reward_list = task.get('action_reward_list', [])
            step_num = task.get('step_num', 0)
            
            if action_annos:
                # 复用标签收集逻辑
                good_steps = (
                    action_annos.get('key_good_before_bad_in_fail', []) +
                    action_annos.get('key_good_in_success', []) +
                    action_annos.get('key_good_after_bad_in_success', [])
                )
                bad_steps = (
                    action_annos.get('key_bad_in_fail', []) +
                    action_annos.get('key_bad_in_success', [])
                )
                
                # 收集该task_type的标签
                for step in good_steps:
                    if 1 <= step <= step_num:
                        task_type_labels[task_type]['true_labels'].append(1)
                        task_type_labels[task_type]['pred_labels'].append(action_reward_list[step - 1])
                
                for step in bad_steps:
                    if 1 <= step <= step_num:
                        task_type_labels[task_type]['true_labels'].append(0)
                        task_type_labels[task_type]['pred_labels'].append(action_reward_list[step - 1])
        
        # 计算按task_type的action reward指标
        task_type_action_metrics = {}
        for task_type, labels in task_type_labels.items():
            if labels['true_labels']:  # 确保有数据
                # 为每个task_type单独筛选结果数据
                task_type_results = [task for task in all_results if task['task_type'] == task_type]
                
                task_type_action_metrics[task_type] = {
                    'metrics': compute_binary_classification_metrics_with_none(
                        labels['true_labels'], 
                        labels['pred_labels']
                    ),
                    # 为每个task_type单独计算group metrics
                    'group_metrics': analyze_group_performance_with_none(task_type_results)
                }
        
        new_stats['action_reward_metrics'] = {
            'overall': overall_action_metrics,
            'by_task_type': task_type_action_metrics,
            'by_reward_type': analyze_group_performance_with_none(all_results)
        }

    # 4. 保存结果
    # 创建results目录（如果不存在）
    results_dir = "results"
    if not os.path.exists(results_dir):
        os.makedirs(results_dir)
    
    # 将metrics文件保存到results目录下
    metrics_filepath = os.path.join(results_dir, metrics_filename)
    with open(metrics_filepath, "w", encoding="utf-8") as f:
        json.dump(new_stats, f, indent=2, ensure_ascii=False)

    # 5. 生成Excel文件
    try:
        # 从metrics文件名生成Excel文件名
        excel_filename = metrics_filename.replace('.json', '.xlsx')
        excel_filepath = os.path.join(results_dir, excel_filename)
        
        # 从投票文件名中提取模型名称
        # 投票文件名格式: voted_{vote_mode}_{eval_mode}_{model_name}.json
        model_name = metrics_filename.replace('voted_', '').replace('.json', '').replace(f'{args.vote_mode}_', '')
        
        # 创建Excel文件
        create_excel_from_metrics(new_stats, model_name, excel_filepath)
        
    except Exception as e:
        print(f"生成Excel文件时出错: {str(e)}")

    # 6. 打印结果
    if 'trajectory_reward_metrics' in new_stats:
        overall = new_stats['trajectory_reward_metrics']['overall']
        print("\nTrajectory Reward Prediction Metrics (Voted):")
        print(f"  Precision: {overall['precision']:.4f}")
        print(f"  Recall: {overall['recall']:.4f}")
        print(f"  Specificity: {overall['specificity']:.4f}")  # 新增specificity显示
        print(f"  NPV: {overall['npv']:.4f}")
        print(f"  F1-score: {overall['f1']:.4f}")
        print(f"  Overall Accuracy (with None): {overall['oa_with_none']:.4f}")
        print(f"  Predictable Accuracy: {overall['oa_predictable']:.4f}")
        print(f"  Coverage: {overall['coverage']:.4f}")
        print(f"  TP: {overall['tp']}, FP: {overall['fp']}, FN: {overall['fn']}, TN: {overall['tn']}")
        print(f"  None: {overall['none_count']} (Positive: {overall['none_positive']}, Negative: {overall['none_negative']})")
        print(f"  Total: {overall['total_samples']}, Predictable: {overall['predictable_samples']}")

    # 打印action reward指标
    if 'action_reward_metrics' in new_stats:
        ar_metrics = new_stats['action_reward_metrics']['overall']
        print("\nAction Reward Prediction Metrics (Voted):")
        print(f"  Precision: {ar_metrics['precision']:.4f}")
        print(f"  Recall: {ar_metrics['recall']:.4f}")
        print(f"  F1-score: {ar_metrics['f1']:.4f}")
        print(f"  Overall Accuracy: {ar_metrics['oa']:.4f}")
        print(f"  Total Samples: {ar_metrics['total_samples']}")
    
    # 打印子类别指标
    if 'action_reward_metrics' in new_stats and 'by_reward_type' in new_stats['action_reward_metrics']:
        group_metrics = new_stats['action_reward_metrics']['by_reward_type']
        print("\nAction Reward Prediction Group Metrics (Voted):")
        for group, metrics in group_metrics.items():
            print(f"\n{group.replace('_', ' ').title()}:")
            print(f"  Accuracy: {metrics['accuracy']:.4f}")
            print(f"  Proportion: {metrics['proportion']:.4f}")
            print(f"  Correct: {metrics['correct']}/{metrics['total']}")
    
    # 打印按task_type的action reward指标
    if 'action_reward_metrics' in new_stats and 'by_task_type' in new_stats['action_reward_metrics']:
        task_type_ar_metrics = new_stats['action_reward_metrics']['by_task_type']
        print("\nAction Reward Prediction Metrics by Task Type (Voted):")
        for task_type, task_metrics in task_type_ar_metrics.items():
            metrics = task_metrics['metrics']
            print(f"\n{task_type.upper()}:")
            print(f"  Precision: {metrics['precision']:.4f}")
            print(f"  Recall: {metrics['recall']:.4f}")
            print(f"  F1-score: {metrics['f1']:.4f}")
            print(f"  Overall Accuracy: {metrics['oa']:.4f}")
            print(f"  Total Samples: {metrics['total_samples']}")
            
            # 打印该task_type的子类别指标
            group_metrics = task_metrics['group_metrics']
            print(f"  Group Metrics:")
            for group, group_stats in group_metrics.items():
                if group_stats['total'] > 0:
                    print(f"    {group.replace('_', ' ').title()}: {group_stats['accuracy']:.4f} ({group_stats['correct']}/{group_stats['total']})")

    print(f"投票指标结果保存在: {metrics_filepath}")
    if 'excel_filepath' in locals():
        print(f"Excel文件保存在: {excel_filepath}")

# ---------- 主流程 ----------
def generate_voted_filename(detailed_filepath_list, vote_mode):
    """从详细结果文件路径列表生成投票结果文件名"""
    model_parts = []
    eval_mode = None
    
    for filepath in detailed_filepath_list:
        filename = os.path.basename(filepath)
        # 解析文件名格式: detailed_results_evalmode_{eval_mode}_cuarewardbench-v0.4_{model_name}.json
        if filename.startswith("detailed_results_evalmode_"):
            parts = filename.replace("detailed_results_evalmode_", "").replace(".json", "").split("_")
            if len(parts) >= 3:
                # 提取评估模式（第一个部分）
                if eval_mode is None:
                    eval_mode = parts[0]
                
                # 提取模型名（从cuarewardbench-v0.4之后的部分）
                model_name_parts = []
                found_benchmark = False
                for part in parts:
                    if found_benchmark:
                        model_name_parts.append(part)
                    elif part.startswith("cuarewardbench"):
                        found_benchmark = True
                
                if model_name_parts:
                    model_name = "_".join(model_name_parts)
                    model_parts.append(f"{eval_mode}_{model_name}")
    
    # 组合最终文件名
    if model_parts:
        models_str = "_".join(model_parts)
        return f"voted_{vote_mode}_{models_str}.json"
    else:
        # 如果解析失败，使用原来的逻辑作为备选
        return "voted_" + os.path.basename(detailed_filepath_list[0]).replace(".json", "_metrics.json")

metrics_filename = generate_voted_filename(detailed_filepath_list, args.vote_mode)

compute_and_save_voting_metrics(
    all_results=merged_results,
    stats=stats,
    task_type_stats=task_type_stats,
    metrics_filename=metrics_filename,
    args=args,
    all_rm_results=all_rm_results
)

print("投票指标已重新计算并保存！")
